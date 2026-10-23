# -*- coding: utf-8 -*-
"""
Script Python tự động tạo đề trắc nghiệm Toán 12 có Bảng Biến Thiên chuẩn Microsoft Word.
Sử dụng thư viện `python-docx` với định dạng bảng, viền và ký hiệu toán học Unicode.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


# ==============================================================================
# HÀM TRỢ GIÚP ĐỊNH DẠNG XML WORD CHO BẢNG VÀ Ô
# ==============================================================================

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    """
    Thiết lập lề trong (padding) cho từng ô tính theo twips (1 pt = 20 twips).
    Giúp nội dung bảng thoáng đãng, không bị dính sát mép viền.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """
    Thiết lập đường viền chính xác cho từng ô bằng XML OpenXML.
    Mỗi cạnh nhận một dict:
        {"val": "single" | "nil" | "dotted" | "dashed", "sz": "4", "color": "000000"}
    Nếu None, giữ nguyên hoặc mặc định.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)

    borders_spec = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for edge, spec in borders_spec.items():
        if spec is not None:
            tag = f'w:{edge}'
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for attr, val in spec.items():
                element.set(qn(f'w:{attr}'), str(val))


def apply_cell_text(cell, text, is_italic=False, is_bold=False, font_size=12):
    """
    Ghi văn bản vào ô, chuẩn hóa font 'Times New Roman' 12pt và căn giữa 2 chiều.
    """
    cell.text = ""  # Xóa paragraph mặc định rỗng nếu có
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0

    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.italic = is_italic
    run.bold = is_bold

    # Căn giữa theo chiều dọc
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


# ==============================================================================
# HÀM DỰNG BẢNG BIẾN THIÊN TỔNG QUÁT (REUSABLE GENERATOR)
# ==============================================================================

def create_variation_table(
    doc: Document,
    x_row: list,
    y_prime_row: list,
    y_top_row: list,
    y_arrow_row: list,
    y_bottom_row: list,
    col_widths: list = None,
    internal_y_style: str = "nil"  # "nil" (ẩn viền) hoặc "dotted" (chấm mảnh)
):
    """
    Tạo Bảng biến thiên chuẩn 5 hàng Toán 12 trong tài liệu docx.

    Tham số:
    - doc: Document của python-docx.
    - x_row: Danh sách giá trị dòng x (không gồm cột tiêu đề).
    - y_prime_row: Danh sách giá trị dòng y'.
    - y_top_row: Tầng trên của y (cực đại, +∞, tiệm cận ngang trên).
    - y_arrow_row: Dòng mũi tên nối tầng trên và tầng dưới (↗, ↘).
    - y_bottom_row: Tầng dưới của y (cực tiểu, -∞, tiệm cận ngang dưới).
    - col_widths: Danh sách độ rộng (Inches) cho từng cột [cột 0, cột 1, ...].
    - internal_y_style: Kiểu viền phân cách nội bộ tầng y ('nil' = ẩn hoàn toàn, 'dotted' = nét đứt).
    """
    num_data_cols = max(len(x_row), len(y_prime_row), len(y_top_row), len(y_arrow_row), len(y_bottom_row))
    total_cols = num_data_cols + 1  # Cột 0 dành cho tiêu đề (x, y', y)

    # Đảm bảo các hàng dữ liệu có đủ số lượng phần tử
    def pad(row_list):
        return row_list + [""] * (num_data_cols - len(row_list))

    x_data = pad(x_row)
    yp_data = pad(y_prime_row)
    yt_data = pad(y_top_row)
    ya_data = pad(y_arrow_row)
    yb_data = pad(y_bottom_row)

    # Khởi tạo bảng 5 hàng
    table = doc.add_table(rows=5, cols=total_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Xóa viền mặc định của Word table để tự kiểm soát viền bằng XML ô
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

    # 1. Điền nhãn cho cột 0
    apply_cell_text(table.cell(0, 0), "x", is_italic=True, is_bold=True)
    apply_cell_text(table.cell(1, 0), "y'", is_italic=True, is_bold=True)

    # Gộp 3 ô hàng 3, 4, 5 (index 2, 3, 4) của cột 0 thành 1 ô duy nhất cho "y"
    cell_y = table.cell(2, 0)
    cell_y.merge(table.cell(4, 0))
    apply_cell_text(cell_y, "y", is_italic=True, is_bold=True)

    # 2. Điền dữ liệu các cột
    for j in range(num_data_cols):
        col_idx = j + 1
        apply_cell_text(table.cell(0, col_idx), str(x_data[j]))
        apply_cell_text(table.cell(1, col_idx), str(yp_data[j]))
        apply_cell_text(table.cell(2, col_idx), str(yt_data[j]))
        apply_cell_text(table.cell(3, col_idx), str(ya_data[j]))
        apply_cell_text(table.cell(4, col_idx), str(yb_data[j]))

    # 3. Tính toán và áp dụng độ rộng cột (Column Widths) cố định
    if col_widths is None or len(col_widths) != total_cols:
        # Tự tính toán phân bổ đều trên chiều ngang trang chuẩn (~5.8 inches)
        header_w = Inches(0.7)
        remain_w = (Inches(5.8) - header_w) / num_data_cols
        widths = [header_w] + [remain_w] * num_data_cols
    else:
        widths = col_widths

    for row in table.rows:
        for i, w in enumerate(widths):
            cell = row.cells[i]
            cell.width = w
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(int(w.twips)))
            tcW.set(qn('w:type'), 'dxa')

    # Đồng bộ thông số tblGrid cho Word
    tblGrid = table._tbl.tblGrid
    if tblGrid is not None:
        tblGrid.clear()
    else:
        tblGrid = OxmlElement('w:tblGrid')
        table._tbl.insert(0, tblGrid)
    for w in widths:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(int(w.twips)))
        tblGrid.append(gc)

    # 4. Định dạng Viền (Borders) thẩm mỹ cao:
    # - Nét đơn đen chuẩn (sz="6" tương đương 0.75pt)
    solid_border = {"val": "single", "sz": "6", "space": "0", "color": "000000"}
    nil_border = {"val": "nil"}
    dotted_border = {"val": "dotted", "sz": "4", "space": "0", "color": "999999"} if internal_y_style == "dotted" else nil_border

    for r_idx in range(5):
        for c_idx in range(total_cols):
            cell = table.cell(r_idx, c_idx)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

            # Xác định viền 4 cạnh của từng ô:
            # - Viền trên: Hàng 0 có viền trên cùng bảng; Hàng 3, 4 có viền nội bộ y (nil/dotted); Hàng 2 có viền phân cách y' và y.
            top_b = solid_border if r_idx == 0 else (solid_border if r_idx == 2 else dotted_border)
            
            # - Viền dưới: Hàng 0 (dưới x) nét đơn; Hàng 1 (dưới y') nét đơn; Hàng 4 (đáy bảng) nét đơn; Hàng 2, 3 nét đứt/ẩn.
            bottom_b = solid_border if r_idx in [0, 1, 4] else dotted_border
            
            # - Viền trái: Cột 0 có viền ngoài cùng bảng; Cột > 0 không có viền dọc ngăn cách.
            left_b = solid_border if c_idx == 0 else nil_border
            
            # - Viền phải: Cột 0 có viền dọc ngăn cách cột tiêu đề; Cột cuối có viền biên bảng.
            right_b = solid_border if c_idx in [0, total_cols - 1] else nil_border

            # Riêng cột 0 đã gộp hàng 2, 3, 4:
            if c_idx == 0 and r_idx in [2, 3, 4]:
                # Viền trên của ô 'y' là đường phân cách dưới hàng 1; viền dưới là đáy bảng
                set_cell_borders(cell, top=solid_border, bottom=solid_border, left=solid_border, right=solid_border)
            else:
                set_cell_borders(cell, top=top_b, bottom=bottom_b, left=left_b, right=right_b)

    # Thêm khoảng trống sau bảng để không dính sát văn bản bên dưới
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(4)

    return table


# ==============================================================================
# HÀM THIẾT LẬP TÀI LIỆU VÀ CÂU HỎI TRẮC NGHIỆM MẪU
# ==============================================================================

def add_multiple_choice_options(doc, opt_a, opt_b, opt_c, opt_d):
    """
    Tạo 4 phương án lựa chọn A, B, C, D dàn hàng ngang cân đối bằng bảng không viền.
    """
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    col_w = Inches(1.5)

    options = [opt_a, opt_b, opt_c, opt_d]
    for i, opt in enumerate(options):
        cell = table.cell(0, i)
        cell.width = col_w
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(opt)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        # Ẩn hoàn toàn viền
        set_cell_borders(cell, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})

    # Đoạn cách sau câu hỏi
    doc.add_paragraph()


def build_sample_docx(output_filename="Cau_Hoi_Trac_Nghiem_Bang_Bien_Thien.docx"):
    """
    Tạo tệp Word chứa các câu hỏi trắc nghiệm Toán 12 với nhiều dạng hàm số khác nhau.
    """
    doc = Document()

    # Cấu hình lề trang tiêu chuẩn (1 inch = 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Cài đặt style Normal font Times New Roman 12pt
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(4)

    # Tiêu đề tài liệu
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BỘ CÂU HỎI TRẮC NGHIỆM TOÁN 12: BẢNG BIẾN THIÊN HÀM SỐ")
    r_title.bold = True
    r_title.font.size = Pt(14)
    r_title.font.name = "Times New Roman"

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("(Chuẩn hóa định dạng bảng biến thiên 5 hàng & Ký hiệu Toán học Unicode)")
    r_sub.italic = True
    r_sub.font.size = Pt(11)
    r_sub.font.name = "Times New Roman"
    doc.add_paragraph()

    # --------------------------------------------------------------------------
    # VÍ DỤ 1: HÀM SỐ BẬC BA y = ax³ + bx² + cx + d (2 điểm cực trị)
    # --------------------------------------------------------------------------
    p_q1 = doc.add_paragraph()
    r_q1_num = p_q1.add_run("Câu 1. ")
    r_q1_num.bold = True
    p_q1.add_run("Cho hàm số bậc ba ")
    r_func = p_q1.add_run("y = f(x)")
    r_func.italic = True
    p_q1.add_run(" có bảng biến thiên như hình dưới đây:")

    # Dữ liệu bảng biến thiên hàm bậc 3: x = -∞, 0, 2, +∞
    # Cột 1 (mốc -∞), Cột 2 (khoảng), Cột 3 (mốc 0), Cột 4 (khoảng), Cột 5 (mốc 2), Cột 6 (khoảng), Cột 7 (mốc +∞)
    x1       = ["−∞",  "",   "0",   "",   "2",   "",   "+∞"]
    yp1      = ["",    "+",  "0",   "−",  "0",   "+",  ""]
    y_top1   = ["",    "",   "3",   "",   "",    "",   "+∞"]
    y_arr1   = ["",    "↗",  "",    "↘",  "",    "↗",  ""]
    y_bot1   = ["−∞",  "",   "",    "",   "−1",  "",   ""]

    # Thiết lập độ rộng cột cân đối (cột nhãn rộng 0.7", các cột mốc/khoảng rộng 0.8")
    widths1 = [Inches(0.7)] + [Inches(0.8)] * 7

    create_variation_table(
        doc,
        x_row=x1,
        y_prime_row=yp1,
        y_top_row=y_top1,
        y_arrow_row=y_arr1,
        y_bottom_row=y_bot1,
        col_widths=widths1,
        internal_y_style="nil"  # Ẩn viền giữa hàng 3, 4, 5
    )

    p_q1_ask = doc.add_paragraph("Hàm số đã cho đồng biến trên khoảng nào dưới đây?")
    add_multiple_choice_options(doc, "A. (0; 2)", "B. (−∞; 0)", "C. (−1; 3)", "D. (0; +∞)")

    # --------------------------------------------------------------------------
    # VÍ DỤ 2: HÀM PHÂN THỨC HỮU TỈ BẬC NHẤT / BẬC NHẤT y = (ax+b)/(cx+d)
    # Có điểm gián đoạn x = 1 (tiệm cận đứng ||, tiệm cận ngang y = 2)
    # --------------------------------------------------------------------------
    p_q2 = doc.add_paragraph()
    r_q2_num = p_q2.add_run("Câu 2. ")
    r_q2_num.bold = True
    p_q2.add_run("Cho hàm số ")
    r_func2 = p_q2.add_run("y = f(x)")
    r_func2.italic = True
    p_q2.add_run(" xác định trên ")
    p_q2.add_run("ℝ \\ {1}")
    p_q2.add_run(" và có bảng biến thiên như sau:")

    # Phân thức nghịch biến trên từng khoảng:
    # x: -∞ -> 1 -> +∞
    # Tại x = 1: y' là ||, tầng y cả trên, giữa và dưới đều thể hiện ký hiệu không xác định ||
    x2       = ["−∞",  "",   "1",   "",   "+∞"]
    yp2      = ["",    "−",  "||",  "−",  ""]
    y_top2   = ["2",   "",   "||",  "",   ""]
    y_arr2   = ["",    "↘",  "||",  "↘",  ""]
    y_bot2   = ["",    "",   "||",  "",   "2"]

    # Tầng dưới cận trái x->1⁻ tiến về −∞, tầng trên cận phải x->1⁺ từ +∞
    # Ta có thể bố trí tinh tế hơn với cột cận:
    x2_b     = ["−∞",  "",    "1",    "",    "+∞"]
    yp2_b    = ["",    "−",   "||",   "−",   ""]
    y_top2_b = ["2",   "",    "+∞",   "",    ""]
    y_arr2_b = ["",    "↘",   "||",   "↘",   ""]
    y_bot2_b = ["",    "−∞",  "||",   "",    "2"]

    widths2 = [Inches(0.7), Inches(0.9), Inches(1.1), Inches(1.0), Inches(1.1), Inches(0.9)]

    create_variation_table(
        doc,
        x_row=x2_b,
        y_prime_row=yp2_b,
        y_top_row=y_top2_b,
        y_arrow_row=y_arr2_b,
        y_bottom_row=y_bot2_b,
        col_widths=widths2,
        internal_y_style="nil"
    )

    p_q2_ask = doc.add_paragraph("Đồ thị hàm số đã cho có đường tiệm cận đứng là:")
    add_multiple_choice_options(doc, "A. x = 1", "B. y = 2", "C. x = 2", "D. y = 1")

    # --------------------------------------------------------------------------
    # VÍ DỤ 3: HÀM BẬC BỐN TRÙNG PHƯƠNG y = ax⁴ + bx² + c (3 điểm cực trị)
    # --------------------------------------------------------------------------
    p_q3 = doc.add_paragraph()
    r_q3_num = p_q3.add_run("Câu 3. ")
    r_q3_num.bold = True
    p_q3.add_run("Cho hàm số ")
    r_func3 = p_q3.add_run("y = f(x)")
    r_func3.italic = True
    p_q3.add_run(" liên tục trên ℝ và có bảng biến thiên như sau:")

    # Hàm trùng phương: cực đại tại x = 0, cực tiểu tại x = ±1
    x3       = ["−∞",  "",   "−1",  "",   "0",   "",   "1",   "",   "+∞"]
    yp3      = ["",    "−",  "0",   "+",  "0",   "−",  "0",   "+",  ""]
    y_top3   = ["+∞",  "",   "",    "",   "2",   "",   "",    "",   "+∞"]
    y_arr3   = ["",    "↘",  "",    "↗",  "",    "↘",  "",    "↗",  ""]
    y_bot3   = ["",    "",   "−3",  "",   "",    "",   "−3",  "",   ""]

    widths3 = [Inches(0.7)] + [Inches(0.6)] * 9

    create_variation_table(
        doc,
        x_row=x3,
        y_prime_row=yp3,
        y_top_row=y_top3,
        y_arrow_row=y_arr3,
        y_bottom_row=y_bot3,
        col_widths=widths3,
        internal_y_style="nil"
    )

    p_q3_ask = doc.add_paragraph("Giá trị cực tiểu của hàm số đã cho bằng:")
    add_multiple_choice_options(doc, "A. −1", "B. 2", "C. −3", "D. 1")

    # Lưu tệp
    doc.save(output_filename)
    print(f"==> Đã tạo thành công file Word tại: {output_filename}")


if __name__ == "__main__":
    build_sample_docx()
